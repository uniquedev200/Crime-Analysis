"""Gate 1 unit tests for XLSX tabular reader.

Marked with @pytest.mark.gate1.
Tests hidden sheets, merged cell handling (no false duplicates), formula caching warnings,
malformed workbook rejection, numeric text formatting, and large workbook limits.
"""

from __future__ import annotations

import datetime
from pathlib import Path

import openpyxl
import pytest

from atlas.contracts.enums import EventType, SheetState
from atlas.contracts.ids import make_doc_id
from atlas.errors import RejectedInputError
from atlas.tabular.cdr import parse_cdr
from atlas.tabular.xlsx_reader import read_xlsx

PINNED_DATE = datetime.datetime(2020, 1, 1, tzinfo=datetime.UTC)


def pin_workbook_properties(wb: openpyxl.Workbook) -> None:
    """Pin workbook created/modified properties for determinism."""
    wb.properties.created = PINNED_DATE
    wb.properties.modified = PINNED_DATE


@pytest.mark.gate1
def test_xlsx_hidden_sheet_flagged(tmp_path: Path) -> None:
    xlsx_file = tmp_path / "sheets_visibility.xlsx"
    wb = openpyxl.Workbook()
    pin_workbook_properties(wb)

    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "VisibleSheet"
    ws1.append(["col1", "col2"])
    ws1.append(["v1", "v2"])

    ws2 = wb.create_sheet(title="HiddenSheet")
    ws2.sheet_state = "hidden"
    ws2.append(["hcol1", "hcol2"])
    ws2.append(["hv1", "hv2"])

    ws3 = wb.create_sheet(title="VeryHiddenSheet")
    ws3.sheet_state = "veryHidden"
    ws3.append(["vhcol1", "vhcol2"])
    ws3.append(["vhv1", "vhv2"])

    wb.save(xlsx_file)
    wb.close()

    result = read_xlsx(xlsx_file)
    sheet_states = {s["name"]: s["state"] for s in result["sheets"]}
    assert sheet_states["VisibleSheet"] == SheetState.VISIBLE.value
    assert sheet_states["HiddenSheet"] == SheetState.HIDDEN.value
    assert sheet_states["VeryHiddenSheet"] == SheetState.VERY_HIDDEN.value


@pytest.mark.gate1
def test_xlsx_merged_range_no_false_duplicates(tmp_path: Path) -> None:
    xlsx_file = tmp_path / "merged.xlsx"
    wb = openpyxl.Workbook()
    pin_workbook_properties(wb)

    ws = wb.active
    assert ws is not None
    ws.append(["a_party", "b_party", "duration_s"])
    ws.append(["9876543210", "9123456780", 60])  # row 2
    ws.append(["", "9123456781", 120])  # row 3

    # Merge a_party across row 2 and 3 (A2:A3)
    ws.merge_cells("A2:A3")
    wb.save(xlsx_file)
    wb.close()

    result = read_xlsx(xlsx_file)
    sheet = result["sheets"][0]
    assert "A2:A3" in sheet["merged_ranges"]

    row2 = sheet["rows"][0]
    row3 = sheet["rows"][1]

    # Row 2 (top-left) keeps the value
    assert row2["a_party"] == "9876543210"
    # Row 3 (other merged cell) is null with in_merged_range=True, preventing false duplication
    assert row3["a_party"] is None


@pytest.mark.gate1
def test_xlsx_formula_without_cache_flagged(tmp_path: Path) -> None:
    xlsx_file = tmp_path / "formula_no_cache.xlsx"
    wb = openpyxl.Workbook()
    pin_workbook_properties(wb)

    ws = wb.active
    assert ws is not None
    ws.append(["val1", "val2", "sum_val"])
    ws.append([10, 20, "=A2+B2"])

    wb.save(xlsx_file)
    wb.close()

    result = read_xlsx(xlsx_file)
    # When openpyxl writes a formula without Excel evaluating it, data_only=True returns None
    assert any("formula_missing_cache" in w for w in result["warnings"])
    row = result["sheets"][0]["rows"][0]
    # Never evaluate formulas: value must be None
    assert row["sum_val"] is None


@pytest.mark.gate1
def test_xlsx_malformed_workbook_handling(tmp_path: Path) -> None:
    bad_file = tmp_path / "corrupt.xlsx"
    # Write garbage bytes that cannot be unzipped
    bad_file.write_bytes(b"PK\x03\x04GARBAGE_NOT_A_VALID_XLSX_STREAM_DATA")

    with pytest.raises(RejectedInputError) as exc_info:
        read_xlsx(bad_file)
    assert "Malformed workbook" in str(exc_info.value)


@pytest.mark.gate1
def test_xlsx_imei_and_leading_zeros(tmp_path: Path) -> None:
    xlsx_file = tmp_path / "imei_leading_zeros.xlsx"
    wb = openpyxl.Workbook()
    pin_workbook_properties(wb)

    ws = wb.active
    assert ws is not None
    ws.append(["imei", "phone", "account"])

    # 15-digit IMEI as integer
    ws.cell(row=2, column=1, value=860123456789012)
    # Phone as text preserving leading zero
    ws.cell(row=2, column=2, value="09876543210")
    # Account with number format requiring 10 digits
    c3 = ws.cell(row=2, column=3, value=12345)
    c3.number_format = "0000000000"

    wb.save(xlsx_file)
    wb.close()

    result = read_xlsx(xlsx_file)
    row = result["sheets"][0]["rows"][0]

    # IMEI never float
    assert row["imei"] == "860123456789012"
    assert "e" not in str(row["imei"]).lower()
    # Phone leading zero preserved
    assert row["phone"] == "09876543210"
    # Number format leading zero applied
    assert row["account"] == "0000012345"


@pytest.mark.gate1
def test_xlsx_large_file_readonly_threshold(tmp_path: Path) -> None:
    xlsx_file = tmp_path / "dummy_large.xlsx"
    wb = openpyxl.Workbook()
    pin_workbook_properties(wb)

    ws = wb.active
    assert ws is not None
    ws.append(["h1", "h2"])
    ws.append(["v1", "v2"])
    wb.save(xlsx_file)
    wb.close()

    # Trigger read-only mode by passing threshold of 0 MB
    result = read_xlsx(xlsx_file, max_full_load_mb=0)
    assert any("merged_cells_not_inspected" in w for w in result["warnings"])
    assert len(result["sheets"]) == 1
    assert len(result["sheets"][0]["rows"]) == 1


@pytest.mark.gate1
def test_xlsx_cdr_parse_with_aggregates(tmp_path: Path) -> None:
    xlsx_file = tmp_path / "cdr.xlsx"
    wb = openpyxl.Workbook()
    pin_workbook_properties(wb)

    ws = wb.active
    assert ws is not None
    ws.append(["a_party", "b_party", "start_time", "duration_s", "call_type"])
    ws.append(["9876543210", "9123456780", datetime.datetime(2023, 1, 1, 10, 0, 0), 60, "VOICE"])
    ws.append(["9876543210", "9123456780", datetime.datetime(2023, 1, 1, 11, 0, 0), 120, "VOICE"])
    ws.append(["9123456780", "9876543210", datetime.datetime(2023, 1, 1, 12, 0, 0), 45, "VOICE"])

    wb.save(xlsx_file)
    wb.close()

    result = read_xlsx(xlsx_file)
    rows = result["sheets"][0]["rows"]

    profile = {
        "name": "cdr_v1",
        "date_format": "%Y-%m-%d %H:%M:%S",
        "timezone": "Asia/Kolkata",
    }
    case_id = "CASE_XLSX_001"
    doc_id = make_doc_id("d" * 64)

    events, aggregates, _mentions = parse_cdr(rows, profile, case_id, doc_id)
    assert len(events) == 3
    # Check invariant: aggregate counts equal event row count
    assert sum(agg.count for agg in aggregates) == len(events)
    assert all(ev.type == EventType.CALL for ev in events)
