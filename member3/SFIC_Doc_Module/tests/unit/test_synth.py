"""Unit tests for synthetic dataset generator, noise models, names, and eval metrics.

Gate 1 test suite validating:
- Deterministic generation across independent runs
- Identifier validity (phone, IMEI, vehicle)
- Ground truth quote provenance against document pages
- Manifest SHA-256 integrity
- Synthetic test data watermarking on all documents
"""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

import openpyxl
import phonenumbers
import pypdfium2 as pdfium
import pytest
from PIL import Image
from synth.generate import generate_all
from synth.names import (
    ALL_PERSONAS,
    SYNTHETIC_LABEL,
    get_case_a_persons,
    get_case_b_persons,
    get_person_by_name,
    list_all_persons,
)
from synth.noise import (
    add_salt_and_pepper,
    add_slight_blur,
    add_slight_rotation,
    apply_digit_occlusion_notch,
    apply_scan_noise,
    substitute_confusable_digit,
)

from atlas.contracts.model_schema import validate_model_output
from atlas.eval.calibrate_ocr import calibrate_ocr, compute_char_level_accuracy
from atlas.eval.metrics import (
    compute_precision_recall_f1,
    eval_exact_match_accuracy,
    eval_mentions_precision_recall,
    eval_ocr_id_review_rate,
    eval_quote_rejection_rate,
)
from atlas.eval.runner import EvalRunner

VEHICLE_REGEX: re.Pattern[str] = re.compile(r"^[A-Z]{2}\d{1,2}[A-Z]{0,3}\d{4}$|^\d{2}BH\d{4}[A-Z]{1,2}$")


def _luhn_checksum(imei: str) -> bool:
    """Validate 15-digit IMEI with Luhn check digit algorithm."""
    if len(imei) != 15 or not imei.isdigit():
        return False
    digits = [int(c) for c in imei]
    total = 0
    for idx, d in enumerate(digits):
        if idx % 2 == 1:
            doubled = d * 2
            total += doubled // 10 + doubled % 10
        else:
            total += d
    return total % 10 == 0


@pytest.mark.gate1
def test_generate_deterministic(tmp_path_factory: pytest.TempPathFactory) -> None:
    """Verify that same seed generates byte-for-byte identical manifest and files across two runs."""
    dir1 = tmp_path_factory.mktemp("synth_run1")
    dir2 = tmp_path_factory.mktemp("synth_run2")

    manifest1 = generate_all(dir1, seed=42)
    manifest2 = generate_all(dir2, seed=42)

    data1 = json.loads(manifest1.read_text(encoding="utf-8"))
    data2 = json.loads(manifest2.read_text(encoding="utf-8"))

    assert data1 == data2, "Synthetic manifests differ between deterministic runs!"

    files1 = {f.name: f for f in dir1.glob("*") if f.is_file()}
    files2 = {f.name: f for f in dir2.glob("*") if f.is_file()}

    assert set(files1.keys()) == set(files2.keys())

    for filename in sorted(files1.keys()):
        bytes1 = files1[filename].read_bytes()
        bytes2 = files2[filename].read_bytes()
        h1 = hashlib.sha256(bytes1).hexdigest()
        h2 = hashlib.sha256(bytes2).hexdigest()
        assert h1 == h2, f"Byte hash mismatch on file '{filename}'!"


@pytest.mark.gate1
def test_identifiers_valid(tmp_path: Path) -> None:
    """Verify every phone, IMEI, and vehicle in truth labels passes its domain validator."""
    generate_all(tmp_path, seed=42)

    truth_files = list(tmp_path.glob("*.truth.json"))
    assert len(truth_files) >= 9, "Expected at least 9 truth files generated"

    checked_phones = 0
    checked_imeis = 0
    checked_vehicles = 0

    for tf in truth_files:
        truth = json.loads(tf.read_text(encoding="utf-8"))
        for mention in truth.get("mentions", []):
            mtype = mention.get("type")
            val = str(mention.get("normalized_value") or mention.get("surface_text", "")).strip()

            if mtype == "PHONE":
                parsed = phonenumbers.parse(val, "IN")
                assert phonenumbers.is_valid_number(parsed), f"Invalid phone in {tf.name}: {val}"
                checked_phones += 1

            elif mtype == "IMEI":
                assert _luhn_checksum(val), f"Invalid Luhn IMEI in {tf.name}: {val}"
                checked_imeis += 1

            elif mtype == "VEHICLE":
                assert VEHICLE_REGEX.match(val), f"Invalid vehicle plate in {tf.name}: {val}"
                checked_vehicles += 1

    assert checked_phones > 0, "No phone numbers found in truth labels"
    assert checked_imeis > 0, "No IMEIs found in truth labels"
    assert checked_vehicles > 0, "No vehicles found in truth labels"


@pytest.mark.gate1
def test_truth_quotes_exist(tmp_path: Path) -> None:
    """Verify every quote in truth labels appears verbatim on its cited document page/text."""
    generate_all(tmp_path, seed=42)

    truth_files = list(tmp_path.glob("*.truth.json"))
    quotes_verified = 0

    for tf in truth_files:
        truth = json.loads(tf.read_text(encoding="utf-8"))
        pages = truth.get("pages", {})

        # Check mention quotes
        for mention in truth.get("mentions", []):
            quote = mention.get("quote")
            page = mention.get("page")
            surface = mention.get("surface_text")

            if quote and page is not None and str(page) in pages:
                page_text = pages[str(page)]
                assert quote in page_text, f"Quote not found on page {page} of {tf.name}: {quote!r}"
                assert surface in quote, f"Mention surface '{surface}' not in quote '{quote}'"
                quotes_verified += 1

            for attr in mention.get("attributes", []):
                attr_quote = attr.get("quote")
                attr_val = attr.get("value_text")
                if attr_quote and page is not None and str(page) in pages:
                    page_text = pages[str(page)]
                    assert attr_quote in page_text, f"Attr quote not in page {page} of {tf.name}: {attr_quote!r}"
                    assert attr_val in attr_quote, f"Attr val '{attr_val}' not in quote '{attr_quote}'"
                    quotes_verified += 1

        # Check relationship quotes
        for rel in truth.get("relationships", []):
            quote = rel.get("quote")
            page = rel.get("page")
            if quote and page is not None and str(page) in pages:
                assert quote in pages[str(page)], f"Rel quote not on page {page} of {tf.name}: {quote!r}"
                quotes_verified += 1

        # Check event quotes
        for ev in truth.get("events", []):
            quote = ev.get("quote")
            page = ev.get("page")
            if quote and page is not None and str(page) in pages:
                assert quote in pages[str(page)], f"Event quote not on page {page} of {tf.name}: {quote!r}"
                quotes_verified += 1

    # Also verify against actual digital PDF text layer for IO_REPORT_A.pdf
    io_pdf = tmp_path / "IO_REPORT_A.pdf"
    assert io_pdf.is_file()
    pdf_doc = pdfium.PdfDocument(io_pdf.read_bytes())
    p2_text = pdf_doc[1].get_textpage().get_text_range()
    assert "Vikrum Singh" in p2_text
    assert "DL01AB1234" in p2_text

    assert quotes_verified > 0, "No quotes were verified!"


@pytest.mark.gate1
def test_manifest_hashes_match(tmp_path: Path) -> None:
    """Verify all manifest SHA-256 entries match the actual generated file bytes."""
    manifest_file = generate_all(tmp_path, seed=42)
    manifest = json.loads(manifest_file.read_text(encoding="utf-8"))

    assert len(manifest) >= 18, f"Manifest should contain at least 18 entries, got {len(manifest)}"

    for sha256_hash, filename in manifest.items():
        file_path = tmp_path / filename
        assert file_path.is_file(), f"File from manifest '{filename}' does not exist!"
        computed_sha = hashlib.sha256(file_path.read_bytes()).hexdigest().lower()
        assert computed_sha == sha256_hash.lower(), f"Hash mismatch for {filename}: {computed_sha} != {sha256_hash}"


@pytest.mark.gate1
def test_synthetic_footer(tmp_path: Path) -> None:
    """Verify every generated document carries the 'SYNTHETIC TEST DATA' label."""
    generate_all(tmp_path, seed=42)

    # 1. Plain text file
    surv_file = tmp_path / "SURVEILLANCE_A.txt"
    assert SYNTHETIC_LABEL in surv_file.read_text(encoding="utf-8")

    # 2. CSV files
    for csv_name in ("BANK_A.csv", "BANK_A_ambiguous.csv", "UNKNOWN_HEADERS.csv"):
        csv_file = tmp_path / csv_name
        assert SYNTHETIC_LABEL in csv_file.read_text(encoding="utf-8"), f"{csv_name} missing synthetic label"

    # 3. Excel workbooks
    for xlsx_name in ("CDR_A.xlsx", "CDR_B.xlsx"):
        xlsx_file = tmp_path / xlsx_name
        wb = openpyxl.load_workbook(xlsx_file, data_only=True)
        found = False
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if cell and SYNTHETIC_LABEL in str(cell):
                        found = True
                        break
                if found:
                    break
            if found:
                break
        assert found, f"{xlsx_name} does not contain '{SYNTHETIC_LABEL}'"

    # 4. Digital PDF
    io_pdf = tmp_path / "IO_REPORT_A.pdf"
    doc = pdfium.PdfDocument(io_pdf.read_bytes())
    for idx, page in enumerate(doc):
        text = page.get_textpage().get_text_range()
        assert SYNTHETIC_LABEL in text, f"IO_REPORT_A.pdf page {idx + 1} missing synthetic label"

    # 5. Image-only PDFs (verify ground truth page source contains watermark)
    for truth_name in ("FIR_A.truth.json", "FIR_B.truth.json"):
        truth_file = tmp_path / truth_name
        truth_data = json.loads(truth_file.read_text(encoding="utf-8"))
        for p_num, p_text in truth_data.get("pages", {}).items():
            assert SYNTHETIC_LABEL in p_text, f"{truth_name} page {p_num} text missing watermark"


@pytest.mark.gate1
def test_names_and_personas() -> None:
    """Verify names.py contains 20+ realistic fictional personas and valid identifiers."""
    persons = list_all_persons()
    assert len(persons) >= 20, f"Expected at least 20 personas, found {len(persons)}"

    case_a = get_case_a_persons()
    assert len(case_a) == 3
    case_b = get_case_b_persons()
    assert len(case_b) == 4

    # Name lookup
    p_vikram = get_person_by_name("Vikram Singh")
    assert p_vikram is not None
    assert "Vicky" in p_vikram.aliases
    assert "Vikrum Singh" in p_vikram.name_variants

    p_vicky = get_person_by_name("Vicky")
    assert p_vicky is not None

    for p in ALL_PERSONAS:
        assert p.label == SYNTHETIC_LABEL
        assert p.age > 0
        assert p.name
        if p.phone:
            parsed = phonenumbers.parse(p.phone, "IN")
            assert phonenumbers.is_valid_number(parsed)
        if p.imei:
            assert _luhn_checksum(p.imei)
        if p.vehicle:
            assert VEHICLE_REGEX.match(p.vehicle)


@pytest.mark.gate1
def test_noise_functions() -> None:
    """Verify noise utilities: rotation, salt-pepper, blur, and confusable digit substitution."""
    im = Image.new("RGB", (200, 200), color=(255, 255, 255))

    rot = add_slight_rotation(im, angle_degrees=0.8)
    assert rot.size == im.size

    sp1 = add_salt_and_pepper(im, amount=0.005, seed=99)
    sp2 = add_salt_and_pepper(im, amount=0.005, seed=99)
    assert list(sp1.getdata()) == list(sp2.getdata()), "Salt and pepper must be deterministic"

    blurred = add_slight_blur(im, radius=0.5)
    assert blurred.size == im.size

    scan1 = apply_scan_noise(im, seed=77)
    scan2 = apply_scan_noise(im, seed=77)
    assert list(scan1.getdata()) == list(scan2.getdata())

    # Confusable digit substitution
    orig = "9876584321"
    conf = substitute_confusable_digit(orig, from_digit="8", to_digit="3", occurrence_index=0)
    assert conf == "9376584321"

    notched = apply_digit_occlusion_notch(im, 50, 50, size=2)
    assert notched.size == im.size


@pytest.mark.gate1
def test_eval_metrics_and_runner(tmp_path: Path) -> None:
    """Verify precision/recall/F1 calculation and EvalRunner evaluation."""
    scores = compute_precision_recall_f1(tp=8, fp=2, fn=2)
    assert scores.precision == 0.8
    assert scores.recall == 0.8
    assert scores.f1 == 0.8

    em = eval_exact_match_accuracy(["+919876543210", "DL01AB1234"], ["+919876543210", "DL01AB1234"])
    assert em == 1.0

    rej_rate = eval_quote_rejection_rate(10, 2)
    assert rej_rate == 0.2

    review_rate = eval_ocr_id_review_rate(5, 5)
    assert review_rate == 1.0

    # Test mentions precision recall
    preds = [
        {"type": "PERSON", "normalized_value": "Vikram Singh"},
        {"type": "PHONE", "normalized_value": "+919876543210"},
        {"type": "PERSON", "normalized_value": "Unknown Accused"},
    ]
    truth = [
        {"type": "PERSON", "normalized_value": "Vikram Singh"},
        {"type": "PHONE", "normalized_value": "+919876543210"},
    ]
    m_eval = eval_mentions_precision_recall(preds, truth)
    assert m_eval["overall"]["tp"] == 2
    assert m_eval["overall"]["fp"] == 1
    assert m_eval["overall"]["fn"] == 0

    # Test EvalRunner with truth directory
    generate_all(tmp_path, seed=42)
    runner = EvalRunner(truth_dir=tmp_path)
    case_a_truth = runner.load_truth_for_case("CASE-A-001")
    assert len(case_a_truth["mentions"]) > 0

    dummy_result = {
        "case_id": "CASE-A-001",
        "mentions": case_a_truth["mentions"][:5],
        "relationships": case_a_truth["relationships"][:2],
        "events": case_a_truth["events"][:1],
        "rejections": [],
    }
    report = runner.evaluate_case(dummy_result, case_a_truth)
    assert report["case_id"] == "CASE-A-001"
    assert "mentions" in report
    assert report["identifier_exact_match_accuracy"] >= 0.0


@pytest.mark.gate1
def test_calibrate_ocr() -> None:
    """Verify OCR confidence calibration and character accuracy calculation."""
    matches, total = compute_char_level_accuracy("9876543210", "9876543210")
    assert matches == 10
    assert total == 10

    sample_tokens = [
        {"ocr_token": "+919876543210", "truth_token": "+919876543210", "confidence": 96.0},
        {"ocr_token": "+919876543210", "truth_token": "+919876543210", "confidence": 98.0},
        {"ocr_token": "DL01AB1234", "truth_token": "DL01AB1234", "confidence": 88.0},
        {"ocr_token": "860123045678900", "truth_token": "860123045678900", "confidence": 92.0},
        {"ocr_token": "9876534210", "truth_token": "98765843210", "confidence": 45.0},
    ]

    res = calibrate_ocr(sample_tokens)
    assert res["status"] == "calibrated"
    assert "buckets" in res
    assert "95_100" in res["buckets"]
    assert res["buckets"]["95_100"]["char_accuracy"] == 1.0


@pytest.mark.gate1
def test_mock_responses_valid() -> None:
    """Verify all scripted mock model responses pass strict ModelOutput schema validation."""
    mock_dir = Path("tests/fixtures/model_responses")
    assert mock_dir.is_dir()

    json_files = list(mock_dir.glob("*.json"))
    assert len(json_files) >= 6, "Expected at least 6 scripted mock responses"

    for jf in json_files:
        content = json.loads(jf.read_text(encoding="utf-8"))
        validated = validate_model_output(content)
        assert "mentions" in validated
        assert "relationships" in validated
        assert "events" in validated
        assert "notes" in validated


@pytest.mark.gate1
def test_synthetic_dataset_fixture(synthetic_dataset: Path) -> None:
    """Verify session fixture synthetic_dataset works from tests/conftest.py."""
    assert synthetic_dataset.is_file()
    assert synthetic_dataset.name == "synthetic_manifest.json"
    content = json.loads(synthetic_dataset.read_text(encoding="utf-8"))
    assert len(content) >= 18
