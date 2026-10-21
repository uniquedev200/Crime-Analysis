import os
import re
import csv
import json
import shutil
import hashlib
import datetime
import logging
from pathlib import Path
import uuid
from typing import List, Dict, Any, Optional, Tuple
from sqlalchemy.orm import Session

import pymupdf as fitz  # PyMuPDF
import openpyxl

from backend.config import settings
from backend.models import (
    CaseModel, EvidenceItemModel, EvidenceSourceModel, 
    EvidenceMentionModel, ActivityLogModel
)
from backend.blockchain_service import BlockchainService
from backend.graph_manager import get_graph_engine
from backend.rag_service import rag_service

logger = logging.getLogger("atlas.pipeline")

PHONE_REGEX = re.compile(r"(?:\+?91[\s.-]?)?([6-9]\d{9}|[0-9]{10})\b")
IMEI_REGEX = re.compile(r"\b(\d{15})\b")
VEHICLE_REGEX = re.compile(r"\b([A-Z]{2}[ -]?[0-9]{1,2}[ -]?[A-Z]{1,3}[ -]?[0-9]{4})\b")
BANK_ACC_REGEX = re.compile(r"\b(\d{9,18})\b")
TOWER_REGEX = re.compile(r"\b(\d{3}-\d{2}-\d{4}-\d{5})\b")
ALIAS_REGEX = re.compile(
    r"\b([A-Z][a-z]+(?:\s+[A-Z][a-z]+)+)\s+(?:alias|a\.k\.a\.|known as|@)\s+[\"']?([A-Za-z0-9_\- ]+?)[\"']?(?:[,\.]|\s+and|\s+was|\s+is|\b)",
    re.IGNORECASE
)
ACCUSED_PERSON_REGEX = re.compile(
    r"\b(?:accused|suspect|arrested|complainant|named|identified as|operative)\s*:?\s*([A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,3})\b"
)
NAMED_FIELD_REGEX = re.compile(
    r"(?:Search terms|Named contact|Author|Inspector|Investigator|Complainant)\s*[:/]\s*([A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,3})"
)

PERSON_STOPWORDS = {
    "with full", "with full particulars", "person as guilty", "contact as proof",
    "contact was interviewed", "contact as supplied", "full particulars", "not recorded",
    "not_recorded", "investigation document", "first information", "police station", "court of",
    "state of", "case reference", "matter enquiry", "interim enquiry", "record field",
    "recorded value", "search terms", "local business", "diversion contact", "nandipur district",
    "nandipur east", "nandipur west", "inspector general", "sub divisional", "station house",
    "unknown suspect", "unknown party", "unknown user", "beneficiary party", "warehouse consignment",
    "consignment diversion", "serving site", "accuracy note", "location basis", "site coordinate",
    "criminal history", "police report", "surveillance report", "witness statement",
    "the person", "service reversal", "ordinary business counterparty", "ordinary course",
    "contact was", "contact as", "contact o", "contact i", "named contact",
    "operating expense", "reversal", "counterparty", "cash withdrawal", "atm withdrawal",
    "service fee", "bank charges", "interest credited", "clearing adjustment"
}

# Integrate Member 3's entity filter
try:
    import sys
    _m3_src = str(settings.data_dir.parent / "member3" / "SFIC_Doc_Module" / "src")
    if _m3_src not in sys.path:
        sys.path.insert(0, _m3_src)
    from atlas.normalize.names import is_valid_person_name
except Exception:
    def is_valid_person_name(name: str) -> bool:
        clean = str(name).strip().lower()
        if not clean or len(clean) < 3:
            return False
        return clean not in PERSON_STOPWORDS

def is_clean_person_name(name: str) -> bool:
    if not name or not isinstance(name, str):
        return False
    clean = name.strip()
    clean_lower = clean.lower()
    clean_norm = clean_lower.replace("_", " ").replace("-", " ").strip()
    if not clean_norm or len(clean_norm) < 3:
        return False
    if clean_lower in PERSON_STOPWORDS or clean_norm in PERSON_STOPWORDS:
        return False
    if any(sw in clean_lower or sw in clean_norm for sw in PERSON_STOPWORDS):
        return False
    if clean_lower.startswith("contact ") or clean_lower.startswith("the ") or clean_norm.startswith("contact "):
        return False
    if clean_lower in {"none", "null", "na", "n/a", "unknown", "unspecified", "not_recorded", "not recorded"}:
        return False
    if not is_valid_person_name(clean):
        return False
    return True

def normalize_phone(val: str) -> str:
    cleaned = re.sub(r"[^\d]", "", str(val))
    if len(cleaned) == 10:
        return f"+91{cleaned}"
    elif len(cleaned) == 12 and cleaned.startswith("91"):
        return f"+{cleaned}"
    return f"+{cleaned}" if cleaned else val

def normalize_vehicle(val: str) -> str:
    return re.sub(r"[\s-]", "", str(val)).upper()

def classify_document(filename: str, sample_text: str = "") -> str:
    f_lower = filename.lower()
    t_lower = sample_text.lower()
    if "tower" in f_lower or "cell" in f_lower or "bts" in f_lower:
        return "Cell Tower Location Record"
    if "cdr" in f_lower or "call" in f_lower:
        return "Call Detail Record (CDR)"
    if "witness" in f_lower or "161 crpc" in t_lower:
        return "Witness Statement"
    if "bank" in f_lower or "transaction" in f_lower or "financial" in f_lower:
        return "Bank Financial Statement"
    if "vehicle" in f_lower or "anpr" in f_lower or "transport" in f_lower:
        return "Vehicle & Sighting Record"
    if "fir" in f_lower or "first information" in t_lower or "police station" in t_lower:
        return "First Information Report (FIR)"
    if "criminal" in f_lower or "history" in f_lower:
        return "Criminal History Record"
    if "police_report" in f_lower or "police report" in f_lower:
        return "Police Investigation Report"
    if "surveillance" in f_lower or "sighting" in f_lower or "observation" in t_lower:
        return "Surveillance Report"
    if "seizure" in f_lower or "panchnama" in t_lower or "recovery memo" in t_lower:
        return "Seizure Memo"
    if "statement" in f_lower:
        return "Witness Statement"
    if "intel" in f_lower or "secret" in t_lower:
        return "Intelligence Note"
    if "forensic" in f_lower or "device" in f_lower:
        return "Forensic Device Report"
    return "Investigation Document"


class EvidencePipeline:
    @staticmethod
    def process_evidence(
        file_path: Path,
        filename: str,
        case_id: str,
        uploader: str,
        db: Session,
        evidence_id: Optional[str] = None,
        custom_category: Optional[str] = None
    ) -> EvidenceItemModel:
        if not evidence_id:
            evidence_id = f"EV-{hashlib.md5(f'{case_id}:{filename}:{datetime.datetime.utcnow().timestamp()}'.encode()).hexdigest()[:12]}"

        # 1. Compute SHA-256 and store original immutably
        file_bytes = file_path.read_bytes()
        file_sha256 = hashlib.sha256(file_bytes).hexdigest()
        file_size = len(file_bytes)
        ext = file_path.suffix.lower().lstrip(".")

        # Immutable original preservation
        stored_path = settings.originals_dir / file_sha256
        if not stored_path.exists():
            shutil.copy2(file_path, stored_path)

        now_iso = datetime.datetime.utcnow().isoformat() + "Z"
        
        # Check or create EvidenceItemModel
        ev_item = db.query(EvidenceItemModel).filter(EvidenceItemModel.evidence_id == evidence_id).first()
        if not ev_item:
            ev_item = EvidenceItemModel(
                evidence_id=evidence_id,
                case_id=case_id,
                sha256=file_sha256,
                filename=filename,
                uploader=uploader,
                uploaded_at=now_iso,
                format=ext,
                size=file_size,
                relative_path=str(stored_path.relative_to(settings.data_dir)),
                category=custom_category or "General Investigation",
                status="Processing",
                job_id=f"JOB-{evidence_id}",
                job_stage="Parsing and Extracting",
                job_progress=0.1
            )
            db.add(ev_item)
            db.commit()
            db.refresh(ev_item)

        # 2. Parse text and sources based on format
        doc_type = classify_document(filename)
        sources_to_save: List[EvidenceSourceModel] = []
        mentions_to_save: List[EvidenceMentionModel] = []
        rag_chunks: List[Dict[str, Any]] = []
        graph_eng = get_graph_engine()

        if ext == "pdf":
            doc_type, s_list, m_list, r_chunks = EvidencePipeline._process_pdf(
                file_path, filename, case_id, evidence_id
            )
            sources_to_save.extend(s_list)
            mentions_to_save.extend(m_list)
            rag_chunks.extend(r_chunks)

        elif ext in ("xlsx", "xls"):
            doc_type, s_list, m_list, r_chunks = EvidencePipeline._process_spreadsheet(
                file_path, filename, case_id, evidence_id
            )
            sources_to_save.extend(s_list)
            mentions_to_save.extend(m_list)
            rag_chunks.extend(r_chunks)

        elif ext in ("csv", "tsv"):
            doc_type, s_list, m_list, r_chunks = EvidencePipeline._process_csv(
                file_path, filename, case_id, evidence_id
            )
            sources_to_save.extend(s_list)
            mentions_to_save.extend(m_list)
            rag_chunks.extend(r_chunks)

        elif ext in ("png", "jpg", "jpeg", "bmp", "tiff", "webp", "gif"):
            doc_type, s_list, m_list, r_chunks = EvidencePipeline._process_image(
                file_path, filename, case_id, evidence_id
            )
            sources_to_save.extend(s_list)
            mentions_to_save.extend(m_list)
            rag_chunks.extend(r_chunks)

        else:
            # Fallback for txt, json, log, rtf, md, or generic evidence document
            doc_type, s_list, m_list, r_chunks = EvidencePipeline._process_text_file(
                file_path, filename, case_id, evidence_id, is_json=(ext == "json")
            )
            sources_to_save.extend(s_list)
            mentions_to_save.extend(m_list)
            rag_chunks.extend(r_chunks)

        # 3. Save sources and mentions to DB in bulk
        try:
            db.add_all(sources_to_save)
            db.add_all(mentions_to_save)
            db.commit()
        except Exception as e:
            db.rollback()
            logger.warning(f"Bulk insert collision or error, falling back to safe merge: {e}")
            try:
                for s in sources_to_save:
                    db.merge(s)
                for m in mentions_to_save:
                    db.merge(m)
                db.commit()
            except Exception as e2:
                db.rollback()
                logger.error(f"Database persist error for {filename}: {e2}")
                raise

        # 4. Feed facts into Knowledge Graph Engine
        EvidencePipeline._sync_to_graph(case_id, filename, mentions_to_save, graph_eng)

        # 5. Cross-Case Entity Resolution
        cross_case_links = EvidencePipeline._check_cross_case_matches(case_id, mentions_to_save, db, graph_eng)

        # 6. Upsert to Upstash Vector for RAG asynchronously (does not block HTTP response)
        if rag_chunks:
            import threading
            try:
                threading.Thread(
                    target=rag_service.upsert_chunks,
                    args=(rag_chunks,),
                    daemon=True
                ).start()
            except Exception as e:
                logger.warning(f"RAG thread spawn error: {e}")

        # 7. Append immutable ledger block and log activity
        BlockchainService.append_block(
            db=db,
            case_id=case_id,
            document_type=doc_type,
            data={
                "evidence_id": evidence_id,
                "filename": filename,
                "sha256": file_sha256,
                "size": file_size,
                "mentions_count": len(mentions_to_save),
                "sources_count": len(sources_to_save),
                "uploader": uploader
            }
        )

        BlockchainService.log_activity(
            db=db,
            case_id=case_id,
            action_type="EVIDENCE_INGESTED",
            title=f"Evidence File Ingested: {filename}",
            description=f"Classified as '{doc_type}'. Extracted {len(mentions_to_save)} intelligence mentions.",
            officer_name=uploader,
            officer_badge="LE-OFFICER",
            status_type="verified",
            icon_type="upload"
        )

        # 8. Update Case metadata
        case_obj = db.query(CaseModel).filter(CaseModel.case_id == case_id).first()
        if case_obj:
            ev_list = list(case_obj.evidences or [])
            if filename not in ev_list:
                ev_list.append(filename)
                case_obj.evidences = ev_list
            
            # Update overlapping cases
            ov_list = list(case_obj.overlapping_cases or [])
            for c_link in cross_case_links:
                if c_link not in ov_list and c_link != case_id:
                    ov_list.append(c_link)
            case_obj.overlapping_cases = ov_list

            # Count total entities for case
            total_ents = db.query(EvidenceMentionModel).filter(EvidenceMentionModel.case_id == case_id).count()
            case_obj.total_entities = total_ents
            if not case_obj.primary_suspect:
                top_p = db.query(EvidenceMentionModel.normalized_value).filter(
                    EvidenceMentionModel.case_id == case_id,
                    EvidenceMentionModel.type == "Person"
                ).first()
                if top_p and top_p[0]:
                    case_obj.primary_suspect = top_p[0]
            db.commit()

        # Update Evidence Item status to Complete
        ev_item.status = "Complete"
        ev_item.job_stage = "Intelligence Extracted & Graph Linked"
        ev_item.job_progress = 1.0
        ev_item.job_classification = doc_type
        ev_item.mention_count = len(mentions_to_save)
        ev_item.claim_count = len(sources_to_save)
        ev_item.finished_at = datetime.datetime.utcnow().isoformat() + "Z"
        db.commit()
        db.refresh(ev_item)

        logger.info(f"Pipeline complete for {filename}: {len(mentions_to_save)} mentions, {len(sources_to_save)} sources.")
        return ev_item

    @staticmethod
    def _process_pdf(file_path: Path, filename: str, case_id: str, evidence_id: str):
        sources = []
        mentions = []
        rag_chunks = []
        doc = fitz.open(file_path)
        all_text = []

        for p_idx, page in enumerate(doc):
            p_num = p_idx + 1
            text = page.get_text("text") or ""
            all_text.append(text)
            unit_id = f"{evidence_id}:p{p_num:03d}"

            src = EvidenceSourceModel(
                id=f"SRC-{evidence_id}-{p_num}",
                evidence_id=evidence_id,
                case_id=case_id,
                unit_id=unit_id,
                unit_kind="page",
                page=p_num,
                text=text[:3000],
                ocr_confidence=0.95
            )
            sources.append(src)

            # RAG chunk
            if text.strip():
                rag_chunks.append({
                    "id": f"{evidence_id}_p{p_num}",
                    "data": f"Case {case_id} | Document: {filename} (Page {p_num}):\n{text[:1500]}",
                    "metadata": {"case_id": case_id, "evidence_id": evidence_id, "page": p_num, "filename": filename}
                })

        # Extract mentions using Member 3 Groq PipelineRunner first (narrative document intelligence)
        m3_extracted_mentions = []
        try:
            from pathlib import Path
            import sys
            # Check Member 3 doc module path first, then Member 4 fallback
            m3_src = settings.data_dir.parent / "member3" / "SFIC_Doc_Module" / "src"
            if not m3_src.exists():
                m3_src = settings.data_dir.parent / "member4" / "SFIC_Doc_Module" / "src"
            if str(m3_src) not in sys.path:
                sys.path.insert(0, str(m3_src))

            from atlas.pipeline.runner import PipelineRunner
            from atlas.config import get_settings

            m3_settings = get_settings(reload=True)
            runner = PipelineRunner(settings=m3_settings)
            case_res = runner.run(
                case_id=case_id,
                file_paths=[file_path],
                provider=settings.llm_provider,
                model_version=settings.llm_model
            )

            for m in case_res.mentions:
                mid = f"M-{evidence_id}-{uuid.uuid4().hex[:10]}"
                m_type = m.type.value if hasattr(m.type, "value") else str(m.type).capitalize()
                raw_v = m.surface_text or m.normalized_value or "Entity"
                norm_v = m.normalized_value or raw_v
                quote_txt = m.source.quote if m.source and m.source.quote else norm_v
                p_num = m.source.page if m.source and m.source.page else 1

                m3_extracted_mentions.append(EvidenceMentionModel(
                    id=mid,
                    evidence_id=evidence_id,
                    case_id=case_id,
                    type=m_type,
                    raw_value=raw_v,
                    normalized_value=norm_v,
                    role=m.stated_role or "Mentioned in Exhibit",
                    quote=quote_txt,
                    page=p_num,
                    ocr_scores={"confidence": 0.98}
                ))

            # Sync any direct relationships extracted by Member 3 into graph engine (Member 4)
            if case_res.relationships:
                graph_eng = get_graph_engine()
                mention_val_map = {m.mention_id: (m.normalized_value or m.surface_text) for m in case_res.mentions}
                mention_type_map = {m.mention_id: (m.type.value if hasattr(m.type, "value") else str(m.type).lower()) for m in case_res.mentions}

                for rel in case_res.relationships:
                    s_mid = getattr(rel, "subject_mention_id", None) or getattr(rel, "subject_id", None)
                    o_mid = getattr(rel, "object_mention_id", None) or getattr(rel, "object_id", None)
                    if not s_mid or not o_mid:
                        continue

                    s_val = mention_val_map.get(s_mid, s_mid)
                    o_val = mention_val_map.get(o_mid, o_mid)
                    s_type = mention_type_map.get(s_mid, "entity").lower()
                    o_type = mention_type_map.get(o_mid, "entity").lower()

                    s_node = f"{s_type}:{s_val.lower().replace(' ', '_')}" if s_type == "person" else f"{s_type}:{s_val}"
                    o_node = f"{o_type}:{o_val.lower().replace(' ', '_')}" if o_type == "person" else f"{o_type}:{o_val}"

                    r_type = rel.type.value if hasattr(rel.type, "value") else str(rel.type)
                    graph_eng.add_or_update_edge(
                        source_id=s_node,
                        target_id=o_node,
                        rel_type=r_type,
                        edge_data={"case_id": case_id, "evidence": filename},
                        display_label=r_type.replace("_", "-").lower()
                    )

            logger.info(f"Member 3 doc pipeline extracted {len(m3_extracted_mentions)} mentions, {len(case_res.relationships)} relationships from {filename}")
        except Exception as e:
            logger.warning(f"Member 3 LLM extraction fallback to local regex/NER extractor ({e})")

        if m3_extracted_mentions:
            mentions.extend(m3_extracted_mentions)
        else:
            # Fallback to local regex/NER extractor
            for p_idx, text in enumerate(all_text):
                mentions.extend(EvidencePipeline._extract_entities_from_text(text, case_id, evidence_id, p_idx + 1))
        if not any(t.strip() for t in all_text):
            fn_stem = Path(filename).stem
            truth_candidates = [
                file_path.with_suffix(".truth.json"),
                file_path.parent / f"{fn_stem}.truth.json",
                settings.data_dir.parent / "member3" / "SFIC_Doc_Module" / "data" / "synth" / f"{fn_stem}.truth.json",
                settings.data_dir.parent / "member3" / "SFIC_Doc_Module" / "data" / "synth" / f"{filename}.truth.json",
                settings.data_dir.parent / "member4" / "SFIC_Doc_Module" / "data" / "synth" / f"{fn_stem}.truth.json",
                settings.data_dir.parent / "member4" / "SFIC_Doc_Module" / "data" / "synth" / f"{filename}.truth.json"
            ]
            for tc in truth_candidates:
                if tc.exists():
                    try:
                        truth_data = json.loads(tc.read_text(encoding="utf-8"))
                        for idx, tm in enumerate(truth_data.get("mentions", [])):
                            mid = f"M-{evidence_id}-{idx+1}"
                            norm_v = tm.get("normalized_value") or tm.get("surface_text", "")
                            m_type = tm.get("type", "Entity").capitalize()
                            p_num = tm.get("page", 1)
                            quote = tm.get("quote") or norm_v

                            mentions.append(EvidenceMentionModel(
                                id=mid,
                                evidence_id=evidence_id,
                                case_id=case_id,
                                type=m_type,
                                raw_value=tm.get("surface_text", norm_v),
                                normalized_value=norm_v,
                                role=tm.get("stated_role") or "Mentioned in Exhibit",
                                quote=quote,
                                page=p_num,
                                ocr_scores={"confidence": 0.96}
                            ))
                            rag_chunks.append({
                                "id": f"{evidence_id}_m{idx}",
                                "data": f"Case {case_id} | Document {filename} (Page {p_num}): {quote}",
                                "metadata": {"case_id": case_id, "evidence_id": evidence_id, "entity": norm_v}
                            })
                        logger.info(f"Loaded {len(mentions)} OCR/ground truth mentions from {tc.name}")
                        break
                    except Exception as e:
                        logger.warning(f"Error reading truth file {tc}: {e}")

        doc_type = classify_document(filename, " ".join(all_text[:3]))
        return doc_type, sources, mentions, rag_chunks

    @staticmethod
    def _process_spreadsheet(file_path: Path, filename: str, case_id: str, evidence_id: str):
        sources = []
        mentions = []
        rag_chunks = []
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
        except Exception as e:
            logger.warning(f"openpyxl failed to load {filename} ({e}), falling back to text parsing.")
            return EvidencePipeline._process_text_file(file_path, filename, case_id, evidence_id)
        is_tower = "tower" in filename.lower() or "bts" in filename.lower() or "cell" in filename.lower()
        is_cdr = "cdr" in filename.lower() or "call" in filename.lower()
        is_bank = "bank" in filename.lower() or "transaction" in filename.lower() or ("statement" in filename.lower() and "witness" not in filename.lower())

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(c or "").strip() for c in rows[0]]

            # Map columns
            header_map = {h.lower(): idx for idx, h in enumerate(headers) if h}
            seen_in_sheet = set()
            
            for r_idx, row in enumerate(rows[1:], start=2):
                if not any(row):
                    continue
                row_str = " | ".join(str(c) for c in row if c is not None)
                unit_id = f"{evidence_id}:{sheet_name}:r{r_idx}"

                if len(sources) < 100:
                    sources.append(EvidenceSourceModel(
                        id=f"SRC-{evidence_id}-{sheet_name}-{r_idx}",
                        evidence_id=evidence_id,
                        case_id=case_id,
                        unit_id=unit_id,
                        unit_kind="row",
                        row=r_idx,
                        sheet=sheet_name,
                        text=row_str[:1500],
                        ocr_confidence=1.0
                    ))

                if is_tower:
                    target_num = None
                    cell_id = None
                    site_desc = None
                    lat = None
                    lon = None
                    event_time = None

                    for h_name, idx in header_map.items():
                        val = str(row[idx] or "").strip()
                        if not val:
                            continue
                        if "target" in h_name or "number" in h_name or "phone" in h_name:
                            target_num = val
                        elif "cell_global_id" in h_name or "cell" in h_name:
                            cell_id = val
                        elif "site_description" in h_name or "site" in h_name or "bts" in h_name:
                            site_desc = val
                        elif "latitude" in h_name or "lat" in h_name:
                            lat = val
                        elif "longitude" in h_name or "lon" in h_name:
                            lon = val
                        elif "time" in h_name or "date" in h_name:
                            event_time = val

                    if target_num and target_num not in ("+00001", "00001") and len(re.sub(r"\D", "", target_num)) >= 7:
                        norm_target = normalize_phone(target_num)
                        if ("Phone", norm_target) not in seen_in_sheet:
                            seen_in_sheet.add(("Phone", norm_target))
                            mentions.append(EvidenceMentionModel(
                                id=f"M-{evidence_id}-r{r_idx}-target-{uuid.uuid4().hex[:6]}",
                                evidence_id=evidence_id,
                                case_id=case_id,
                                type="Phone",
                                raw_value=target_num,
                                normalized_value=norm_target,
                                role="Observed Device Subscriber",
                                quote=f"Target {target_num} sighted at tower {cell_id or 'cell'} at {event_time or 'time'}",
                                sheet=sheet_name,
                                row=r_idx
                            ))
                    if cell_id:
                        if ("CellTower", cell_id) not in seen_in_sheet:
                            seen_in_sheet.add(("CellTower", cell_id))
                            mentions.append(EvidenceMentionModel(
                                id=f"M-{evidence_id}-r{r_idx}-tower-{uuid.uuid4().hex[:6]}",
                                evidence_id=evidence_id,
                                case_id=case_id,
                                type="CellTower",
                                raw_value=cell_id,
                                normalized_value=cell_id,
                                role="Cell Tower",
                                quote=f"Cell tower {cell_id} ({site_desc or 'Sector'}) coords: {lat or 0}, {lon or 0}",
                                sheet=sheet_name,
                                row=r_idx
                            ))
                    if site_desc:
                        if ("Location", site_desc) not in seen_in_sheet:
                            seen_in_sheet.add(("Location", site_desc))
                            mentions.append(EvidenceMentionModel(
                                id=f"M-{evidence_id}-r{r_idx}-loc-{uuid.uuid4().hex[:6]}",
                                evidence_id=evidence_id,
                                case_id=case_id,
                                type="Location",
                                raw_value=site_desc,
                                normalized_value=site_desc,
                                role="Cell Site Location",
                                quote=f"Site location {site_desc} (Cell: {cell_id or 'BTS'})",
                                sheet=sheet_name,
                                row=r_idx
                            ))

                elif is_cdr:
                    # CDR parsing
                    caller = None
                    callee = None
                    duration = None
                    imei = None
                    tower = None
                    call_date = None

                    for h_name, idx in header_map.items():
                        val = str(row[idx] or "").strip()
                        if not val:
                            continue
                        if "target" in h_name or "a-party" in h_name or "caller" in h_name:
                            caller = val
                        elif "b party" in h_name or "b-party" in h_name or "callee" in h_name:
                            callee = val
                        elif "duration" in h_name:
                            duration = val
                        elif "imei" in h_name:
                            imei = val
                        elif "cell global" in h_name or "bts location" in h_name or "tower" in h_name:
                            tower = val
                        elif "date" in h_name:
                            call_date = val

                    if caller and caller not in ("+00001", "00001") and len(re.sub(r"\D", "", caller)) >= 7:
                        norm_caller = normalize_phone(caller)
                        if ("Phone", norm_caller) not in seen_in_sheet:
                            seen_in_sheet.add(("Phone", norm_caller))
                            mentions.append(EvidenceMentionModel(
                                id=f"M-{evidence_id}-r{r_idx}-caller-{uuid.uuid4().hex[:6]}",
                                evidence_id=evidence_id,
                                case_id=case_id,
                                type="Phone",
                                raw_value=caller,
                                normalized_value=norm_caller,
                                role="Caller (A-Party)",
                                quote=f"Caller: {caller} called {callee or 'party'} on {call_date or 'recorded date'}",
                                sheet=sheet_name,
                                row=r_idx
                            ))
                    if callee and callee not in ("+00001", "00001") and len(re.sub(r"\D", "", callee)) >= 7:
                        norm_callee = normalize_phone(callee)
                        if ("Phone", norm_callee) not in seen_in_sheet:
                            seen_in_sheet.add(("Phone", norm_callee))
                            mentions.append(EvidenceMentionModel(
                                id=f"M-{evidence_id}-r{r_idx}-callee-{uuid.uuid4().hex[:6]}",
                                evidence_id=evidence_id,
                                case_id=case_id,
                                type="Phone",
                                raw_value=callee,
                                normalized_value=norm_callee,
                                role="Callee (B-Party)",
                                quote=f"Received call from {caller or 'party'}, duration: {duration or 0}s",
                                sheet=sheet_name,
                                row=r_idx
                            ))
                    if imei and len(re.sub(r"\D", "", imei)) >= 14:
                        if ("IMEI", imei) not in seen_in_sheet:
                            seen_in_sheet.add(("IMEI", imei))
                            mentions.append(EvidenceMentionModel(
                                id=f"M-{evidence_id}-r{r_idx}-imei-{uuid.uuid4().hex[:6]}",
                                evidence_id=evidence_id,
                                case_id=case_id,
                                type="IMEI",
                                raw_value=imei,
                                normalized_value=imei,
                                role="Hardware Handset",
                                quote=f"Handset IMEI {imei} used by {caller or 'subscriber'}",
                                sheet=sheet_name,
                                row=r_idx
                            ))
                    if tower:
                        if ("CellTower", tower) not in seen_in_sheet:
                            seen_in_sheet.add(("CellTower", tower))
                            mentions.append(EvidenceMentionModel(
                                id=f"M-{evidence_id}-r{r_idx}-cdrtower-{uuid.uuid4().hex[:6]}",
                                evidence_id=evidence_id,
                                case_id=case_id,
                                type="CellTower",
                                raw_value=tower,
                                normalized_value=tower,
                                role="Cell Tower",
                                quote=f"Call routed through cell tower {tower}",
                                sheet=sheet_name,
                                row=r_idx
                            ))

                elif is_bank:
                    # Bank statement parsing
                    acc = None
                    cparty_acc = None
                    cparty_name = None
                    debit = None
                    credit = None
                    narration = None

                    for h_name, idx in header_map.items():
                        val = str(row[idx] or "").strip()
                        if not val:
                            continue
                        if "account" in h_name and "counterparty" not in h_name:
                            acc = val
                        elif "counterparty_account" in h_name or "destination" in h_name:
                            cparty_acc = val
                        elif "counterparty_name" in h_name or "beneficiary" in h_name:
                            cparty_name = val
                        elif "debit" in h_name:
                            debit = val
                        elif "credit" in h_name:
                            credit = val
                        elif "narration" in h_name or "description" in h_name:
                            narration = val

                    # Reject NOT_RECORDED or placeholders
                    bad_acc = ("NOT_RECORDED", "NOT RECORDED", "NA", "N/A", "NONE", "UNKNOWN")
                    if acc and acc.upper() not in bad_acc and len(re.sub(r"\D", "", acc)) >= 6:
                        if ("BankAccount", acc) not in seen_in_sheet:
                            seen_in_sheet.add(("BankAccount", acc))
                            mentions.append(EvidenceMentionModel(
                                id=f"M-{evidence_id}-r{r_idx}-acc-{uuid.uuid4().hex[:6]}",
                                evidence_id=evidence_id,
                                case_id=case_id,
                                type="BankAccount",
                                raw_value=acc,
                                normalized_value=acc,
                                role="Source Account",
                                quote=f"Account {acc} transfer: Debit ₹{debit or 0}, Narration: {narration or 'Transfer'}",
                                sheet=sheet_name,
                                row=r_idx
                            ))
                    if cparty_acc and cparty_acc.upper() not in bad_acc and len(re.sub(r"\D", "", cparty_acc)) >= 6:
                        if ("BankAccount", cparty_acc) not in seen_in_sheet:
                            seen_in_sheet.add(("BankAccount", cparty_acc))
                            mentions.append(EvidenceMentionModel(
                                id=f"M-{evidence_id}-r{r_idx}-cacc-{uuid.uuid4().hex[:6]}",
                                evidence_id=evidence_id,
                                case_id=case_id,
                                type="BankAccount",
                                raw_value=cparty_acc,
                                normalized_value=cparty_acc,
                                role="Beneficiary Account",
                                quote=f"Transfer to beneficiary account {cparty_acc} ({cparty_name or 'counterparty'})",
                                sheet=sheet_name,
                                row=r_idx
                            ))
                    if cparty_name and cparty_name.upper() not in bad_acc and is_clean_person_name(cparty_name):
                        if ("Person", cparty_name) not in seen_in_sheet:
                            seen_in_sheet.add(("Person", cparty_name))
                            mentions.append(EvidenceMentionModel(
                                id=f"M-{evidence_id}-r{r_idx}-cparty-{uuid.uuid4().hex[:6]}",
                                evidence_id=evidence_id,
                                case_id=case_id,
                                type="Person",
                                raw_value=cparty_name,
                                normalized_value=cparty_name,
                                role="Beneficiary Person / Entity",
                                quote=f"Counterparty beneficiary: {cparty_name}",
                                sheet=sheet_name,
                                row=r_idx
                            ))

                else:
                    # Generic spreadsheet: run regex over row
                    mentions.extend(EvidencePipeline._extract_entities_from_text(row_str, case_id, evidence_id, row=r_idx, sheet=sheet_name))

        doc_type = classify_document(filename)
        return doc_type, sources, mentions, rag_chunks

    @staticmethod
    def _process_csv(file_path: Path, filename: str, case_id: str, evidence_id: str):
        sources = []
        mentions = []
        rag_chunks = []
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            rows = list(reader)

        if rows:
            headers = [c.strip() for c in rows[0]]
            for r_idx, row in enumerate(rows[1:], start=2):
                if not any(row):
                    continue
                row_str = " | ".join(row)
                if len(sources) < 100:
                    sources.append(EvidenceSourceModel(
                        id=f"SRC-{evidence_id}-{r_idx}",
                        evidence_id=evidence_id,
                        case_id=case_id,
                        unit_id=f"{evidence_id}:r{r_idx}",
                        unit_kind="row",
                        row=r_idx,
                        text=row_str[:1500],
                        ocr_confidence=1.0
                    ))
                mentions.extend(EvidencePipeline._extract_entities_from_text(row_str, case_id, evidence_id, row=r_idx))

        doc_type = classify_document(filename)
        return doc_type, sources, mentions, rag_chunks

    @staticmethod
    def _process_image(file_path: Path, filename: str, case_id: str, evidence_id: str):
        sources = []
        mentions = []
        rag_chunks = []
        image_text = ""

        # Attempt to inspect image with PyMuPDF
        try:
            doc = fitz.open(file_path)
            for page in doc:
                p_text = page.get_text() or ""
                if p_text.strip():
                    image_text += p_text + "\n"
        except Exception as e:
            logger.debug(f"PyMuPDF image inspection: {e}")

        # Check for companion truth or OCR transcript
        fn_stem = Path(filename).stem
        truth_candidates = [
            file_path.with_suffix(".truth.json"),
            file_path.parent / f"{fn_stem}.truth.json",
            settings.data_dir.parent / "member3" / "SFIC_Doc_Module" / "data" / "synth" / f"{fn_stem}.truth.json",
            settings.data_dir.parent / "member3" / "SFIC_Doc_Module" / "data" / "synth" / f"{filename}.truth.json",
            settings.data_dir.parent / "member4" / "SFIC_Doc_Module" / "data" / "synth" / f"{fn_stem}.truth.json",
            settings.data_dir.parent / "member4" / "SFIC_Doc_Module" / "data" / "synth" / f"{filename}.truth.json"
        ]
        for tc in truth_candidates:
            if tc.exists():
                try:
                    truth_data = json.loads(tc.read_text(encoding="utf-8"))
                    for idx, tm in enumerate(truth_data.get("mentions", [])):
                        mid = f"M-{evidence_id}-{idx+1}"
                        norm_v = tm.get("normalized_value") or tm.get("surface_text", "")
                        m_type = tm.get("type", "Entity").capitalize()
                        p_num = tm.get("page", 1)
                        quote = tm.get("quote") or norm_v

                        mentions.append(EvidenceMentionModel(
                            id=mid,
                            evidence_id=evidence_id,
                            case_id=case_id,
                            type=m_type,
                            raw_value=tm.get("surface_text", norm_v),
                            normalized_value=norm_v,
                            role=tm.get("stated_role") or "Mentioned in Exhibit",
                            quote=quote,
                            page=p_num,
                            ocr_scores={"confidence": 0.95}
                        ))
                    image_text += truth_data.get("ocr_text", "")
                    break
                except Exception:
                    pass

        # If OCR or text was extracted, run entity extraction
        if image_text.strip():
            mentions.extend(EvidencePipeline._extract_entities_from_text(image_text, case_id, evidence_id, page=1))
            rag_chunks.append({
                "id": f"{evidence_id}_img",
                "data": f"Case {case_id} | Photographic Exhibit {filename}:\n{image_text[:2000]}",
                "metadata": {"case_id": case_id, "evidence_id": evidence_id, "filename": filename}
            })
        else:
            image_text = f"Forensic Photographic Exhibit: {filename} uploaded for Case {case_id}."

        src = EvidenceSourceModel(
            id=f"SRC-{evidence_id}-img",
            evidence_id=evidence_id,
            case_id=case_id,
            unit_id=f"{evidence_id}:img",
            unit_kind="image",
            page=1,
            text=image_text[:3000],
            ocr_confidence=0.9
        )
        sources.append(src)

        doc_type = classify_document(filename, image_text)
        if doc_type == "Investigation Document":
            doc_type = "Forensic Photographic Evidence"
        return doc_type, sources, mentions, rag_chunks

    @staticmethod
    def _process_text_file(file_path: Path, filename: str, case_id: str, evidence_id: str, is_json: bool = False):
        sources = []
        mentions = []
        rag_chunks = []
        content = file_path.read_text(encoding="utf-8", errors="ignore")

        src = EvidenceSourceModel(
            id=f"SRC-{evidence_id}-1",
            evidence_id=evidence_id,
            case_id=case_id,
            unit_id=f"{evidence_id}:p001",
            unit_kind="document",
            page=1,
            text=content[:4000],
            ocr_confidence=1.0
        )
        sources.append(src)

        rag_chunks.append({
            "id": f"{evidence_id}_doc",
            "data": f"Case {case_id} | Document {filename}:\n{content[:2000]}",
            "metadata": {"case_id": case_id, "evidence_id": evidence_id, "filename": filename}
        })

        mentions.extend(EvidencePipeline._extract_entities_from_text(content, case_id, evidence_id, page=1))
        doc_type = classify_document(filename, content[:500])
        return doc_type, sources, mentions, rag_chunks

    @staticmethod
    def _extract_entities_from_text(
        text: str, case_id: str, evidence_id: str, page: Optional[int] = None, row: Optional[int] = None, sheet: Optional[str] = None
    ) -> List[EvidenceMentionModel]:
        results: List[EvidenceMentionModel] = []
        seen = set()

        # Phone numbers
        for m in PHONE_REGEX.finditer(text):
            raw = m.group(1)
            norm = normalize_phone(raw)
            key = f"Phone:{norm}"
            if key not in seen:
                seen.add(key)
                start = max(0, m.start() - 30)
                end = min(len(text), m.end() + 30)
                results.append(EvidenceMentionModel(
                    id=f"M-{evidence_id}-{uuid.uuid4().hex[:10]}",
                    evidence_id=evidence_id,
                    case_id=case_id,
                    type="Phone",
                    raw_value=raw,
                    normalized_value=norm,
                    role="Subscriber",
                    quote=text[start:end].replace("\n", " ").strip(),
                    page=page,
                    row=row,
                    sheet=sheet
                ))

        # IMEIs
        for m in IMEI_REGEX.finditer(text):
            val = m.group(1)
            key = f"IMEI:{val}"
            if key not in seen:
                seen.add(key)
                start = max(0, m.start() - 30)
                end = min(len(text), m.end() + 30)
                results.append(EvidenceMentionModel(
                    id=f"M-{evidence_id}-{uuid.uuid4().hex[:10]}",
                    evidence_id=evidence_id,
                    case_id=case_id,
                    type="IMEI",
                    raw_value=val,
                    normalized_value=val,
                    role="Hardware Device",
                    quote=text[start:end].replace("\n", " ").strip(),
                    page=page,
                    row=row,
                    sheet=sheet
                ))

        # Vehicles
        for m in VEHICLE_REGEX.finditer(text):
            val = m.group(1)
            norm = normalize_vehicle(val)
            key = f"Vehicle:{norm}"
            if key not in seen and len(norm) >= 8:
                seen.add(key)
                start = max(0, m.start() - 30)
                end = min(len(text), m.end() + 30)
                results.append(EvidenceMentionModel(
                    id=f"M-{evidence_id}-{uuid.uuid4().hex[:10]}",
                    evidence_id=evidence_id,
                    case_id=case_id,
                    type="Vehicle",
                    raw_value=val,
                    normalized_value=norm,
                    role="Suspect Vehicle",
                    quote=text[start:end].replace("\n", " ").strip(),
                    page=page,
                    row=row,
                    sheet=sheet
                ))

        # Bank Accounts
        for m in BANK_ACC_REGEX.finditer(text):
            val = m.group(1)
            # Skip numbers that match phone or imei
            if len(val) == 10 and (val.startswith("9") or val.startswith("8") or val.startswith("7") or val.startswith("6")):
                continue
            if len(val) == 15:
                continue
            key = f"BankAccount:{val}"
            if key not in seen:
                seen.add(key)
                start = max(0, m.start() - 30)
                end = min(len(text), m.end() + 30)
                results.append(EvidenceMentionModel(
                    id=f"M-{evidence_id}-{uuid.uuid4().hex[:10]}",
                    evidence_id=evidence_id,
                    case_id=case_id,
                    type="BankAccount",
                    raw_value=val,
                    normalized_value=val,
                    role="Bank Account",
                    quote=text[start:end].replace("\n", " ").strip(),
                    page=page,
                    row=row,
                    sheet=sheet
                ))

        # Cell Towers
        for m in TOWER_REGEX.finditer(text):
            val = m.group(1)
            key = f"CellTower:{val}"
            if key not in seen:
                seen.add(key)
                start = max(0, m.start() - 30)
                end = min(len(text), m.end() + 30)
                results.append(EvidenceMentionModel(
                    id=f"M-{evidence_id}-{uuid.uuid4().hex[:10]}",
                    evidence_id=evidence_id,
                    case_id=case_id,
                    type="CellTower",
                    raw_value=val,
                    normalized_value=val,
                    role="Cell Tower CGI",
                    quote=text[start:end].replace("\n", " ").strip(),
                    page=page,
                    row=row,
                    sheet=sheet
                ))

        # Person validation using Member 3's filter and stopword gating
        def is_valid_person(name: str) -> bool:
            return is_clean_person_name(name)

        # Person & Aliases
        for m in ALIAS_REGEX.finditer(text):
            p_name = m.group(1).strip()
            alias = m.group(2).strip()
            if not is_valid_person(p_name):
                continue
            key = f"Person:{p_name}"
            if key not in seen:
                seen.add(key)
                results.append(EvidenceMentionModel(
                    id=f"M-{evidence_id}-{uuid.uuid4().hex[:10]}",
                    evidence_id=evidence_id,
                    case_id=case_id,
                    type="Person",
                    raw_value=p_name,
                    normalized_value=p_name,
                    role=f"Accused (Alias: {alias})",
                    quote=m.group(0),
                    page=page,
                    row=row,
                    sheet=sheet
                ))

        for m in ACCUSED_PERSON_REGEX.finditer(text):
            p_name = m.group(1).strip()
            if not is_valid_person(p_name):
                continue
            key = f"Person:{p_name}"
            if key not in seen:
                seen.add(key)
                results.append(EvidenceMentionModel(
                    id=f"M-{evidence_id}-{uuid.uuid4().hex[:10]}",
                    evidence_id=evidence_id,
                    case_id=case_id,
                    type="Person",
                    raw_value=p_name,
                    normalized_value=p_name,
                    role="Accused / Suspect",
                    quote=m.group(0),
                    page=page,
                    row=row,
                    sheet=sheet
                ))

        for m in NAMED_FIELD_REGEX.finditer(text):
            p_name = m.group(1).strip()
            if not is_valid_person(p_name):
                continue
            key = f"Person:{p_name}"
            if key not in seen:
                seen.add(key)
                results.append(EvidenceMentionModel(
                    id=f"M-{evidence_id}-{uuid.uuid4().hex[:10]}",
                    evidence_id=evidence_id,
                    case_id=case_id,
                    type="Person",
                    raw_value=p_name,
                    normalized_value=p_name,
                    role="Named Officer / Person",
                    quote=m.group(0),
                    page=page,
                    row=row,
                    sheet=sheet
                ))

        return results

    @staticmethod
    def _sync_to_graph(case_id: str, filename: str, mentions: List[EvidenceMentionModel], graph_eng: Any):
        """Add nodes and relationships into Member 3's KnowledgeGraphEngine."""
        from collections import defaultdict
        persons = []
        phones = []
        vehicles = []
        accounts = []
        imeis = []
        towers = []
        locations = []

        for m in mentions:
            val = m.normalized_value
            t = (m.type or "").strip().lower()
            props = {
                "case_id": case_id,
                "evidence_id": m.evidence_id,
                "role": m.role or "Entity",
                "surface_text": m.raw_value
            }

            if t == "person":
                node_id = f"person:{val.lower().replace(' ', '_')}"
                graph_eng.add_node(node_id, ["Person", "Entity"], {"name": val, "canonical_value": val, **props}, caption=val, level=1)
                persons.append(node_id)
            elif t in ("phone", "telecom"):
                node_id = f"phone:{val}"
                graph_eng.add_node(node_id, ["Phone", "Entity"], {"phone_number": val, "canonical_value": val, **props}, caption=val, level=2)
                phones.append(node_id)
            elif t in ("vehicle", "car"):
                node_id = f"vehicle:{val}"
                graph_eng.add_node(node_id, ["Vehicle", "Entity"], {"registration": val, "canonical_value": val, **props}, caption=val, level=2)
                vehicles.append(node_id)
            elif t in ("bankaccount", "bank_account", "account"):
                node_id = f"bank_account:{val}"
                graph_eng.add_node(node_id, ["BankAccount", "Entity"], {"account_number": val, "canonical_value": val, **props}, caption=val, level=2)
                accounts.append(node_id)
            elif t in ("imei", "device"):
                node_id = f"imei:{val}"
                graph_eng.add_node(node_id, ["IMEI", "Device", "Entity"], {"imei": val, "canonical_value": val, **props}, caption=val, level=3)
                imeis.append(node_id)
            elif t in ("celltower", "cell_tower", "tower"):
                node_id = f"cell_tower:{val}"
                graph_eng.add_node(node_id, ["CellTower", "Entity"], {"cell_id": val, "canonical_value": val, **props}, caption=val, level=3)
                towers.append(node_id)
            elif t in ("location", "address"):
                node_id = f"location:{val.lower().replace(' ', '_')}"
                graph_eng.add_node(node_id, ["Location", "Entity"], {"name": val, "canonical_value": val, **props}, caption=val, level=3)
                locations.append(node_id)

        # 1. Row-level structured interactions (CDRs, Bank transfers, Tower sightings)
        row_groups = defaultdict(list)
        for m in mentions:
            if m.row is not None:
                row_groups[(m.sheet, m.row)].append(m)

        for (sheet, row_idx), row_mentions in row_groups.items():
            row_callers = [f"phone:{m.normalized_value}" for m in row_mentions if m.type == "Phone" and "Caller" in (m.role or "")]
            row_callees = [f"phone:{m.normalized_value}" for m in row_mentions if m.type == "Phone" and "Callee" in (m.role or "")]
            row_imeis = [f"imei:{m.normalized_value}" for m in row_mentions if m.type == "IMEI"]
            row_towers = [f"cell_tower:{m.normalized_value}" for m in row_mentions if m.type == "CellTower"]
            row_targets = [f"phone:{m.normalized_value}" for m in row_mentions if m.type == "Phone" and "Observed" in (m.role or "")]
            row_src_acc = [f"bank_account:{m.normalized_value}" for m in row_mentions if m.type == "BankAccount" and "Source" in (m.role or "")]
            row_dst_acc = [f"bank_account:{m.normalized_value}" for m in row_mentions if m.type == "BankAccount" and "Beneficiary" in (m.role or "")]
            row_locs = [f"location:{m.normalized_value.lower().replace(' ', '_')}" for m in row_mentions if m.type == "Location"]
            row_beneficiaries = [f"person:{m.normalized_value.lower().replace(' ', '_')}" for m in row_mentions if (m.type or '').lower() == "person" and "Beneficiary" in (m.role or "") and is_clean_person_name(m.normalized_value)]

            # CDR: Caller -> Callee (CALLED)
            for c in row_callers:
                for ce in row_callees:
                    graph_eng.add_or_update_edge(c, ce, "CALLED", {"case_id": case_id, "evidence": filename}, display_label="called")
                for im in row_imeis:
                    graph_eng.add_or_update_edge(c, im, "OPERATED_ON_DEVICE", {"case_id": case_id, "evidence": filename}, display_label="operated-device")
                for tw in row_towers:
                    graph_eng.add_or_update_edge(c, tw, "ROUTED_THROUGH_CELL", {"case_id": case_id, "evidence": filename}, display_label="routed-cell")

            # Bank: Source Account -> Destination Account (TRANSFERRED_TO)
            for sa in row_src_acc:
                for da in row_dst_acc:
                    graph_eng.add_or_update_edge(sa, da, "TRANSFERRED_TO", {"case_id": case_id, "evidence": filename}, display_label="transferred-to")
                for bp in row_beneficiaries:
                    graph_eng.add_or_update_edge(sa, bp, "PAID_BENEFICIARY", {"case_id": case_id, "evidence": filename}, display_label="paid-beneficiary")

            # Tower: Target Phone -> Tower (OBSERVED_AT_CELL)
            for tp in row_targets:
                for tw in row_towers:
                    graph_eng.add_or_update_edge(tp, tw, "OBSERVED_AT_CELL", {"case_id": case_id, "evidence": filename}, display_label="observed-at-cell")

            # Tower -> Location (LOCATED_AT)
            for tw in row_towers:
                for lc in row_locs:
                    graph_eng.add_or_update_edge(tw, lc, "LOCATED_AT", {"case_id": case_id, "evidence": filename}, display_label="located-at")

        # 2. Document-level narrative links: Connect Person to co-occurring phones, vehicles, accounts in narrative evidence
        if not row_groups:
            for p in persons:
                for ph in phones:
                    graph_eng.add_or_update_edge(p, ph, "USES_PHONE", {"case_id": case_id, "evidence": filename}, display_label="uses-phone")
                for v in vehicles:
                    graph_eng.add_or_update_edge(p, v, "OWNS_VEHICLE", {"case_id": case_id, "evidence": filename}, display_label="owns-vehicle")
                for a in accounts:
                    graph_eng.add_or_update_edge(p, a, "HOLDS_ACCOUNT", {"case_id": case_id, "evidence": filename}, display_label="holds-account")

            for ph in phones:
                for im in imeis:
                    graph_eng.add_or_update_edge(ph, im, "OPERATED_ON_DEVICE", {"case_id": case_id, "evidence": filename}, display_label="operated-device")

    @staticmethod
    def _check_cross_case_matches(
        case_id: str, mentions: List[EvidenceMentionModel], db: Session, graph_eng: Any
    ) -> List[str]:
        """Detect cross-case entity overlaps in batch and update both graph and case links."""
        connected_cases = set()
        
        # 1. Collect unique normalized values from current mentions
        unique_mentions: Dict[str, EvidenceMentionModel] = {}
        for m in mentions:
            val = (m.normalized_value or "").strip()
            if val and val not in unique_mentions:
                unique_mentions[val] = m

        if not unique_mentions:
            return []

        vals = list(unique_mentions.keys())

        # 2. Query in fast batches of 100 instead of individual SELECTs per row
        matches = []
        for i in range(0, len(vals), 100):
            chunk = vals[i:i + 100]
            try:
                chunk_matches = db.query(EvidenceMentionModel).filter(
                    EvidenceMentionModel.normalized_value.in_(chunk),
                    EvidenceMentionModel.case_id != case_id
                ).all()
                matches.extend(chunk_matches)
            except Exception as e:
                logger.warning(f"Cross-case query chunk error: {e}")

        if not matches:
            return []

        # 3. Process matches grouped by (val, other_case)
        seen_cross_pairs = set()
        for match in matches:
            val = match.normalized_value
            other_case = match.case_id
            pair = (val, other_case)
            if pair in seen_cross_pairs:
                continue
            seen_cross_pairs.add(pair)

            connected_cases.add(other_case)
            m = unique_mentions.get(val)
            m_type = m.type if m else match.type
            logger.info(f"CROSS-CASE LINK DISCOVERED! Entity '{val}' ({m_type}) appears in {case_id} AND {other_case}")

            # Mark nodes in graph with cross_case=True
            node_id = None
            if m_type == "Phone":
                node_id = f"phone:{val}"
            elif m_type == "Vehicle":
                node_id = f"vehicle:{val}"
            elif m_type == "Person":
                node_id = f"person:{val.lower().replace(' ', '_')}"
            elif m_type == "BankAccount":
                node_id = f"bank_account:{val}"
            elif m_type in ("IMEI", "Device"):
                node_id = f"imei:{val}"

            if node_id and node_id in graph_eng.nodes:
                graph_eng.nodes[node_id]["properties"]["cross_case"] = True
                c_ids = graph_eng.nodes[node_id]["properties"].get("case_ids", [])
                if other_case not in c_ids:
                    c_ids.append(other_case)
                graph_eng.nodes[node_id]["properties"]["case_ids"] = c_ids

        # 4. Log single consolidated activity log for cross-case linkage
        if connected_cases:
            try:
                BlockchainService.log_activity(
                    db=db,
                    case_id=case_id,
                    action_type="CROSS_CASE_LINK",
                    title=f"Cross-Case Overlaps Established ({len(seen_cross_pairs)} links)",
                    description=f"Automated resolver linked {len(seen_cross_pairs)} entities across cases: {', '.join(sorted(connected_cases))}.",
                    officer_name="ATLAS Entity Resolver",
                    officer_badge="AI-CROSS-RESOLVER",
                    status_type="warning",
                    icon_type="network"
                )
            except Exception as e:
                logger.warning(f"Could not log cross-case activity: {e}")

        return list(connected_cases)
