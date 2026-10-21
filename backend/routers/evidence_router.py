import os
import shutil
import csv
import logging
import hashlib
import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, Form, status
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel
from sqlalchemy.orm import Session
import openpyxl

from backend.config import settings
from backend.database import get_db
from backend.models import (
    CaseModel, EvidenceItemModel, EvidenceSourceModel,
    EvidenceMentionModel, EvidenceReviewModel
)
from backend.pipeline import EvidencePipeline
from backend.blockchain_service import BlockchainService

logger = logging.getLogger("atlas.evidence_router")

router = APIRouter(tags=["Evidence Ingestion & Forensic Processing"])

class ClassificationUpdateRequest(BaseModel):
    category: str
    reviewer: Optional[str] = "Investigator"

class ReviewCorrectionRequest(BaseModel):
    target_type: str = "mentions"
    target_id: str
    decision: str
    correction: Optional[str] = None
    reviewer: Optional[str] = "Investigator"

def serialize_evidence_item(e: EvidenceItemModel) -> Dict[str, Any]:
    return {
        "evidence_id": e.evidence_id,
        "case_id": e.case_id,
        "sha256": e.sha256,
        "filename": e.filename,
        "uploader": e.uploader,
        "uploaded_at": e.uploaded_at,
        "format": e.format,
        "size": e.size,
        "relative_path": e.relative_path,
        "classification_override": e.classification_override,
        "category": e.category,
        "job": {
            "job_id": e.job_id or f"JOB-{e.evidence_id}",
            "status": e.status,
            "stage": e.job_stage or "Completed",
            "progress": e.job_progress or 1.0,
            "created_at": e.uploaded_at,
            "finished_at": e.finished_at,
            "error": e.job_error,
            "warnings": e.job_warnings or [],
            "classification": e.classification_override or e.job_classification or "Document",
            "mention_count": e.mention_count or 0,
            "claim_count": e.claim_count or 0
        }
    }

@router.get("/api/evidence/cases", summary="List Evidence Cases with Count")
async def get_evidence_cases(db: Session = Depends(get_db)):
    distinct_cases = db.query(EvidenceItemModel.case_id).distinct().all()
    items = []
    for (cid,) in distinct_cases:
        count = db.query(EvidenceItemModel).filter(EvidenceItemModel.case_id == cid).count()
        c_obj = db.query(CaseModel).filter(CaseModel.case_id == cid).first()
        items.append({
            "case_id": cid,
            "count": count,
            "case_name": c_obj.case_name if c_obj else cid,
            "status": c_obj.status if c_obj else "Active"
        })
    return {"items": items}

@router.get("/api/evidence", summary="List Evidence Items for Case")
async def list_evidence(
    case_id: Optional[str] = Query(None),
    limit: int = Query(200, ge=1, le=1000),
    db: Session = Depends(get_db)
):
    query = db.query(EvidenceItemModel)
    if case_id and case_id != "ALL":
        query = query.filter(EvidenceItemModel.case_id == case_id)
    items = query.order_by(EvidenceItemModel.created_at.desc()).limit(limit).all()
    return {"items": [serialize_evidence_item(e) for e in items]}

@router.post("/api/evidence/uploads", summary="Upload Case Evidence Files")
@router.post("/api/evidence/upload", summary="Upload Case Evidence Files (Alias)")
async def upload_evidence_files(
    files: Optional[List[UploadFile]] = File(None),
    file: Optional[UploadFile] = File(None),
    case_id: Optional[str] = Form(None),
    case_id_query: Optional[str] = Query(None, alias="case_id"),
    category: Optional[str] = Form("General Investigation"),
    uploader: Optional[str] = Form("Investigator"),
    db: Session = Depends(get_db)
):
    # Consolidate files from either 'files' or 'file'
    file_list: List[UploadFile] = []
    if files:
        file_list.extend(files)
    if file:
        file_list.append(file)

    if not file_list:
        raise HTTPException(status_code=400, detail="No evidence files were attached to the upload request.")

    cid = (case_id or case_id_query or "").strip()
    if not cid:
        cid = f"CASE_{datetime.datetime.utcnow().strftime('%Y%m%d_%H%M%S')}"

    # Ensure case exists in DB
    c_obj = db.query(CaseModel).filter(CaseModel.case_id == cid).first()
    if not c_obj:
        c_obj = CaseModel(
            case_id=cid,
            fir_number=f"FIR-{cid}",
            case_name=f"Case {cid}",
            status="Under Investigation",
            lead_investigator=uploader or "Lead Detective",
            uploaded_by=uploader or "Investigator",
            crime_category=category or "General Investigation"
        )
        db.add(c_obj)
        db.commit()

    uploaded_items = []
    case_upload_dir = settings.uploads_dir / cid
    case_upload_dir.mkdir(parents=True, exist_ok=True)

    for upload_file in file_list:
        fname = upload_file.filename or f"evidence_{datetime.datetime.utcnow().strftime('%H%M%S')}.dat"
        target_path = case_upload_dir / fname
        try:
            with open(target_path, "wb") as f:
                shutil.copyfileobj(upload_file.file, f)

            # Run pipeline
            item = EvidencePipeline.process_evidence(
                file_path=target_path,
                filename=fname,
                case_id=cid,
                uploader=uploader or "Investigator",
                db=db,
                custom_category=category
            )
            uploaded_items.append(serialize_evidence_item(item))
        except Exception as e:
            logger.error(f"Error processing uploaded file {fname}: {e}", exc_info=True)
            db.rollback()
            fail_id = f"EV-{hashlib.md5(f'{cid}:{fname}'.encode()).hexdigest()[:12]}"
            failed_item = EvidenceItemModel(
                evidence_id=fail_id,
                case_id=cid,
                sha256="",
                filename=fname,
                uploader=uploader or "Investigator",
                uploaded_at=datetime.datetime.utcnow().isoformat() + "Z",
                format=Path(fname).suffix.lower().lstrip("."),
                size=0,
                relative_path=str(target_path),
                category=category or "General Investigation",
                status="Failed",
                job_id=f"JOB-{fail_id}",
                job_stage="Processing Failed",
                job_progress=0.0,
                job_error=str(e)
            )
            try:
                db.merge(failed_item)
                db.commit()
            except Exception:
                db.rollback()
            uploaded_items.append(serialize_evidence_item(failed_item))

    return {"items": uploaded_items}

@router.get("/api/evidence/{evidence_id}", summary="Get Evidence Details")
async def get_evidence_detail(evidence_id: str, db: Session = Depends(get_db)):
    item = db.query(EvidenceItemModel).filter(EvidenceItemModel.evidence_id == evidence_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Evidence item not found.")
    return serialize_evidence_item(item)

@router.get("/api/evidence/{evidence_id}/original", summary="Download or Stream Evidence File")
async def get_evidence_original(
    evidence_id: str,
    inline: bool = Query(False),
    db: Session = Depends(get_db)
):
    item = db.query(EvidenceItemModel).filter(EvidenceItemModel.evidence_id == evidence_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Evidence item not found.")

    file_path = settings.originals_dir / item.sha256
    if not file_path.exists():
        # Fallback to uploads dir
        file_path = settings.uploads_dir / item.case_id / item.filename

    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Original evidence file missing from storage.")

    media_type = "application/octet-stream"
    ext = item.format.lower()
    if ext == "pdf":
        media_type = "application/pdf"
    elif ext in ("jpg", "jpeg"):
        media_type = "image/jpeg"
    elif ext == "png":
        media_type = "image/png"
    elif ext in ("csv", "txt"):
        media_type = "text/plain"
    elif ext in ("xlsx", "xls"):
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    disposition = f"{'inline' if inline else 'attachment'}; filename=\"{item.filename}\""
    return FileResponse(
        path=file_path,
        media_type=media_type,
        headers={"Content-Disposition": disposition}
    )

@router.get("/api/evidence/{evidence_id}/results/sources", summary="Extracted Text and Row Sources")
async def get_evidence_sources(
    evidence_id: str,
    limit: int = Query(500, ge=1, le=1000),
    db: Session = Depends(get_db)
):
    sources = db.query(EvidenceSourceModel).filter(
        EvidenceSourceModel.evidence_id == evidence_id
    ).limit(limit).all()
    items = []
    for s in sources:
        items.append({
            "id": s.id,
            "evidence_id": s.evidence_id,
            "case_id": s.case_id,
            "unit_id": s.unit_id,
            "unit_kind": s.unit_kind,
            "page": s.page,
            "row": s.row,
            "sheet": s.sheet,
            "text": s.text,
            "bbox": s.bbox,
            "ocr_confidence": s.ocr_confidence
        })
    return {"items": items}

@router.get("/api/evidence/{evidence_id}/results/mentions", summary="Extracted Entity Mentions & Provenance")
async def get_evidence_mentions(
    evidence_id: str,
    limit: int = Query(500, ge=1, le=1000),
    db: Session = Depends(get_db)
):
    mentions = db.query(EvidenceMentionModel).filter(
        EvidenceMentionModel.evidence_id == evidence_id
    ).limit(limit).all()
    items = []
    for m in mentions:
        items.append({
            "id": m.id,
            "evidence_id": m.evidence_id,
            "case_id": m.case_id,
            "type": m.type,
            "raw_value": m.raw_value,
            "normalized_value": m.normalized_value,
            "role": m.role,
            "identifier_trust": m.identifier_trust,
            "quote": m.quote,
            "quote_status": m.quote_status,
            "provenance": {
                "page": m.page,
                "sheet": m.sheet,
                "row": m.row,
                "snippet": m.quote,
                "bbox": m.bbox
            },
            "ocr": m.ocr_scores,
            "confirmed_by": m.confirmed_by
        })
    return {"items": items}

@router.post("/api/evidence/{evidence_id}/reviewed", summary="Mark Evidence as Reviewed")
async def mark_reviewed(evidence_id: str, db: Session = Depends(get_db)):
    item = db.query(EvidenceItemModel).filter(EvidenceItemModel.evidence_id == evidence_id).first()
    if item:
        item.status = "Reviewed"
        db.commit()
    return {"status": "reviewed", "evidence_id": evidence_id}

@router.post("/api/evidence/{evidence_id}/retry", summary="Reprocess Evidence File")
async def retry_evidence(evidence_id: str, db: Session = Depends(get_db)):
    item = db.query(EvidenceItemModel).filter(EvidenceItemModel.evidence_id == evidence_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Evidence item not found.")
    
    file_path = settings.originals_dir / item.sha256
    if not file_path.exists():
        file_path = settings.uploads_dir / item.case_id / item.filename

    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Evidence file not found on disk.")

    # Re-run pipeline
    EvidencePipeline.process_evidence(
        file_path=file_path,
        filename=item.filename,
        case_id=item.case_id,
        uploader=item.uploader,
        db=db,
        evidence_id=item.evidence_id
    )
    return {"status": "reprocessed", "evidence_id": evidence_id}

@router.post("/api/evidence/{evidence_id}/classification", summary="Update Classification")
async def update_classification(
    evidence_id: str,
    req: ClassificationUpdateRequest,
    db: Session = Depends(get_db)
):
    item = db.query(EvidenceItemModel).filter(EvidenceItemModel.evidence_id == evidence_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Evidence item not found.")
    item.classification_override = req.category
    db.commit()
    return {"status": "updated", "category": req.category}

@router.get("/api/evidence/{evidence_id}/reviews", summary="Get Review Corrections")
async def get_reviews(evidence_id: str, db: Session = Depends(get_db)):
    revs = db.query(EvidenceReviewModel).filter(EvidenceReviewModel.evidence_id == evidence_id).all()
    return {"reviews": revs}

@router.post("/api/evidence/{evidence_id}/reviews", summary="Submit Human Correction")
async def submit_review(
    evidence_id: str,
    req: ReviewCorrectionRequest,
    db: Session = Depends(get_db)
):
    rev = EvidenceReviewModel(
        evidence_id=evidence_id,
        target_type=req.target_type,
        target_id=req.target_id,
        decision=req.decision,
        correction=req.correction,
        reviewer=req.reviewer or "Investigator"
    )
    db.add(rev)

    # If it is a mention correction, update the mention normalized value
    mention = db.query(EvidenceMentionModel).filter(EvidenceMentionModel.id == req.target_id).first()
    if mention and req.correction:
        mention.normalized_value = req.correction
        mention.confirmed_by = req.reviewer or "Investigator"

    db.commit()
    return {"status": "correction_recorded", "target_id": req.target_id}

@router.delete("/api/evidence/{evidence_id}", summary="Delete Evidence Item")
async def delete_evidence(
    evidence_id: str,
    officer_name: Optional[str] = Query("Investigator"),
    officer_badge: Optional[str] = Query("LE-8402"),
    db: Session = Depends(get_db)
):
    item = db.query(EvidenceItemModel).filter(EvidenceItemModel.evidence_id == evidence_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Evidence item not found.")

    cid = item.case_id
    db.query(EvidenceMentionModel).filter(EvidenceMentionModel.evidence_id == evidence_id).delete()
    db.query(EvidenceSourceModel).filter(EvidenceSourceModel.evidence_id == evidence_id).delete()
    db.query(EvidenceReviewModel).filter(EvidenceReviewModel.evidence_id == evidence_id).delete()
    db.delete(item)
    db.commit()

    BlockchainService.log_activity(
        db=db,
        case_id=cid,
        action_type="EVIDENCE_DELETED",
        title=f"Evidence File Purged: {item.filename}",
        description=f"Evidence ID {evidence_id} removed by {officer_name} ({officer_badge}).",
        officer_name=officer_name,
        officer_badge=officer_badge,
        status_type="warning",
        icon_type="trash"
    )

    return {"status": "deleted", "evidence_id": evidence_id}

@router.delete("/api/evidence/cases/{case_id}", summary="Delete All Evidence for Case")
async def delete_case_evidence(case_id: str, db: Session = Depends(get_db)):
    db.query(EvidenceMentionModel).filter(EvidenceMentionModel.case_id == case_id).delete()
    db.query(EvidenceSourceModel).filter(EvidenceSourceModel.case_id == case_id).delete()
    db.query(EvidenceItemModel).filter(EvidenceItemModel.case_id == case_id).delete()
    db.commit()
    return {"status": "all_evidence_deleted", "case_id": case_id}

# ---------- Ingest and Spreadsheet Preview Compatibility Endpoints ----------

@router.post("/api/v1/ingest", summary="Fast Multi-File Ingest Endpoint")
async def v1_ingest(
    case_id: Optional[str] = Query(None),
    case_id_form: Optional[str] = Form(None, alias="case_id"),
    files: Optional[List[UploadFile]] = File(None),
    file: Optional[UploadFile] = File(None),
    uploaded_by: Optional[str] = Query(None),
    officer_name: Optional[str] = Query(None),
    category: Optional[str] = Form("General Investigation"),
    db: Session = Depends(get_db)
):
    # Consolidate files from either 'files' or 'file'
    file_list: List[UploadFile] = []
    if files:
        file_list.extend(files)
    if file:
        file_list.append(file)

    if not file_list:
        raise HTTPException(status_code=400, detail="No evidence files were attached for ingestion.")

    cid = (case_id or case_id_form or "").strip()
    if not cid:
        cid = f"CASE_{datetime.datetime.utcnow().strftime('%Y%m%d_%H%M%S')}"

    c_obj = db.query(CaseModel).filter(CaseModel.case_id == cid).first()
    if not c_obj:
        c_obj = CaseModel(
            case_id=cid,
            fir_number=f"FIR-{cid}",
            case_name=f"Case {cid}",
            status="Under Investigation",
            lead_investigator=officer_name or "Lead Detective",
            uploaded_by=uploaded_by or "Investigator"
        )
        db.add(c_obj)
        db.commit()

    case_upload_dir = settings.uploads_dir / cid
    case_upload_dir.mkdir(parents=True, exist_ok=True)

    processed_count = 0
    for upload_file in file_list:
        fname = upload_file.filename or f"evidence_{datetime.datetime.utcnow().strftime('%H%M%S')}.dat"
        target_path = case_upload_dir / fname
        try:
            with open(target_path, "wb") as f:
                shutil.copyfileobj(upload_file.file, f)

            EvidencePipeline.process_evidence(
                file_path=target_path,
                filename=fname,
                case_id=cid,
                uploader=officer_name or uploaded_by or "Investigator",
                db=db,
                custom_category=category
            )
            processed_count += 1
        except Exception as e:
            logger.error(f"Error during ingest for file {fname}: {e}", exc_info=True)

    return {"case_id": cid, "status": "ingested", "files_count": len(file_list), "processed_count": processed_count}

@router.post("/api/v1/preview", summary="Spreadsheet Preview (Excel / CSV)")
async def preview_spreadsheet(file: UploadFile = File(...)):
    filename = file.filename or "spreadsheet.xlsx"
    ext = Path(filename).suffix.lower()
    content = await file.read()

    sheets = []
    if ext == ".csv":
        text = content.decode("utf-8", errors="ignore")
        reader = csv.reader(text.splitlines())
        rows = list(reader)
        sheets.append({
            "name": "CSV",
            "rows": rows[:200],
            "row_count": len(rows),
            "column_count": max([len(r) for r in rows], default=0)
        })
    else:
        # Excel
        import io
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            raw_rows = list(ws.iter_rows(values_only=True))
            str_rows = [[str(cell if cell is not None else "") for cell in r] for r in raw_rows[:200]]
            col_count = max([len(r) for r in str_rows], default=0)
            sheets.append({
                "name": sname,
                "rows": str_rows,
                "row_count": len(raw_rows),
                "column_count": col_count
            })

    return {"sheets": sheets}
